import io
import re
from datetime import datetime, timedelta, timezone, tzinfo

import bcrypt
import chardet
import jwt
from pypdf import PdfReader

from app.config import settings

_TZ_CN = timezone(timedelta(hours=8))


def now_cn() -> datetime:
    return datetime.now(_TZ_CN)


def hash_password(plain: str) -> str:
    return bcrypt.hashpw(plain.encode("utf-8"), bcrypt.gensalt(rounds=10)).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except ValueError:
        return False


def create_token(user_id: int, username: str, role: str) -> str:
    payload = {
        "sub": str(user_id),
        "username": username,
        "role": role,
        "exp": datetime.now(_TZ_CN) + timedelta(minutes=settings.JWT_EXPIRE_MINUTES),
    }
    return jwt.encode(payload, settings.SECRET_KEY, algorithm=settings.JWT_ALGORITHM)


def decode_token(token: str) -> dict:
    return jwt.decode(token, settings.SECRET_KEY, algorithms=[settings.JWT_ALGORITHM])


def role_from_username(username: str) -> str:
    low = username.lower()
    if low.startswith("admin"):
        return "super_admin"
    if low.startswith("lawyer"):
        return "lawyer"
    return "user"


ROLE_LABELS = {"super_admin": "超级管理员", "lawyer": "律师", "user": "普通用户"}

DOMAIN_KEYWORDS = {
    "合同法": ["合同", "要约", "承诺", "违约", "合同纠纷", "买卖", "租赁", "借款", "定金", "履行", "解除", "格式条款", "缔约", "合同编"],
    "劳动法": ["劳动合同", "用人单位", "劳动者", "试用期", "工资", "加班", "社会保险", "社保", "经济补偿", "裁员", "工伤", "竞业限制", "解雇", "劳动争议"],
    "知识产权法": ["专利", "商标", "著作权", "版权", "知识产权", "许可", "专有权", "发明", "作品", "注册商标", "实用新型", "外观设计"],
    "公司法": ["公司", "股东", "董事会", "股权", "有限责任", "股份", "法定代表人", "注册资本", "清算", "分红", "公司治理", "监事"],
    "婚姻家庭法": ["婚姻", "离婚", "夫妻", "抚养", "继承", "遗嘱", "子女", "赡养", "共同财产", "监护", "亲属"],
    "侵权责任法": ["侵权", "损害赔偿", "过错", "人身损害", "精神损害", "侵害", "安全保障义务", "无过错责任", "公平责任"],
    "刑法": ["犯罪", "刑罚", "有期徒刑", "刑事", "诈骗", "盗窃", "量刑", "自首", "故意伤害", "拘役"],
    "行政法": ["行政", "行政许可", "行政处罚", "行政复议", "行政诉讼", "行政机关", "执法", "听证"],
}

ALL_KEYWORDS = sorted({kw for kws in DOMAIN_KEYWORDS.values() for kw in kws}, key=len, reverse=True)


def classify_domain(text: str) -> str:
    scores = {}
    for domain, kws in DOMAIN_KEYWORDS.items():
        score = sum(text.count(kw) for kw in kws)
        if score > 0:
            scores[domain] = score
    if not scores:
        return "综合法律"
    return max(scores, key=scores.get)


def extract_keywords(text: str, limit: int = 8) -> list:
    counts = {}
    for kw in ALL_KEYWORDS:
        n = text.count(kw)
        if n > 0:
            counts[kw] = n
    ranked = sorted(counts.items(), key=lambda x: (-x[1], -len(x[0])))
    return [k for k, _ in ranked[:limit]]


def chunk_text(text: str, size: int = None, overlap: int = None) -> list:
    size = size or settings.CHUNK_SIZE
    overlap = overlap or settings.CHUNK_OVERLAP
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= size:
        return [text] if text else []
    sentences = re.split(r"(?<=[。！？；!?;\n])", text)
    chunks, cur = [], ""
    for s in sentences:
        if len(cur) + len(s) <= size:
            cur += s
        else:
            if cur.strip():
                chunks.append(cur.strip())
            if len(s) > size:
                for i in range(0, len(s), size - overlap):
                    chunks.append(s[i : i + size].strip())
                cur = ""
            else:
                cur = s
    if cur.strip():
        chunks.append(cur.strip())
    final = []
    for c in chunks:
        if not final or len(final[-1]) + len(c) > size + overlap:
            final.append(c)
        else:
            final[-1] += c
    return final


def decode_bytes(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    guess = chardet.detect(data[:60000])
    enc = (guess or {}).get("encoding") or "utf-8"
    for e in [enc, "utf-8", "gb18030"]:
        try:
            return data.decode(e)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="ignore")


def parse_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def parse_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[工作表: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def parse_uploaded_file(filename: str, data: bytes) -> tuple:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext in ("txt", "md", "markdown", "csv", "json", "html", "xml", "log"):
        text = decode_bytes(data)
    elif ext == "pdf":
        text = parse_pdf(data)
    elif ext == "docx":
        text = parse_docx(data)
    elif ext in ("xlsx", "xlsm"):
        text = parse_xlsx(data)
    elif ext == "doc":
        text = decode_bytes(data)
    else:
        text = decode_bytes(data)
    return ext or "txt", text


def sanitize_filename(name: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|]", "_", name).strip()
    return name[:120] or "file"